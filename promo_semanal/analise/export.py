"""Briefing exportável: Excel (uma aba por seção) e HTML autocontido com os gráficos."""
from __future__ import annotations

import html as html_mod
from io import BytesIO

import pandas as pd
import plotly.graph_objects as go
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def excel(abas: dict[str, pd.DataFrame]) -> bytes:
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for nome, df in abas.items():
            nome = nome[:31]
            df.to_excel(w, sheet_name=nome, index=False)
            ws = w.sheets[nome]
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="184F95")
            for i, col in enumerate(df.columns, 1):
                largura = max([len(str(col))] + [len(str(x)) for x in df[col].head(200)])
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, largura + 2), 60)
            ws.freeze_panes = "A2"
    return buf.getvalue()


CSS = """
:root { --bg:#f9f9f7; --surface:#fcfcfb; --ink:#0b0b0b; --ink2:#52514e; --line:#e1e0d9; --accent:#184f95; }
@media (prefers-color-scheme: dark) { :root:not([data-theme="light"]) {
  --bg:#0d0d0d; --surface:#1a1a19; --ink:#ffffff; --ink2:#c3c2b7; --line:#2c2c2a; --accent:#86b6ef; } }
:root[data-theme="dark"] { --bg:#0d0d0d; --surface:#1a1a19; --ink:#ffffff; --ink2:#c3c2b7; --line:#2c2c2a; --accent:#86b6ef; }
* { box-sizing: border-box; }
body { margin:0; background:var(--bg); color:var(--ink); font-family: system-ui, -apple-system, "Segoe UI", sans-serif; }
main { max-width: 1180px; margin: 0 auto; padding: 24px 16px 64px; }
h1 { font-size: 1.6rem; margin: 0 0 4px; } h2 { font-size: 1.2rem; margin: 32px 0 8px; color: var(--accent); }
h3 { font-size: 1rem; margin: 20px 0 6px; }
p.meta { color: var(--ink2); margin: 0 0 16px; font-size: .9rem; }
section { background: var(--surface); border: 1px solid var(--line); border-radius: 10px; padding: 16px; margin-top: 16px; }
.tabela { overflow-x: auto; }
table { border-collapse: collapse; width: 100%; font-size: .85rem; font-variant-numeric: tabular-nums; }
th, td { padding: 6px 8px; border-bottom: 1px solid var(--line); text-align: right; white-space: nowrap; }
th:first-child, td:first-child { text-align: left; } th { color: var(--ink2); font-weight: 600; }
.nota { color: var(--ink2); font-size: .8rem; }
@media print { section { break-inside: avoid; } }
"""


def _tabela(df: pd.DataFrame) -> str:
    fmt = df.copy()
    for c in fmt.columns:
        if pd.api.types.is_float_dtype(fmt[c]):
            fmt[c] = fmt[c].map(lambda x: "" if pd.isna(x) else f"{x:,.2f}".replace(",", "§").replace(".", ",").replace("§", "."))
    return f'<div class="tabela">{fmt.to_html(index=False, border=0, na_rep="", escape=True)}</div>'


def html(titulo: str, subtitulo: str, secoes: list[tuple[str, list]]) -> str:
    """secoes: [(título, [itens])]; item = DataFrame, Figure, ('h3', texto) ou texto (nota)."""
    partes, primeiro_grafico = [], True
    for nome, itens in secoes:
        corpo = []
        for it in itens:
            if isinstance(it, pd.DataFrame):
                corpo.append(_tabela(it))
            elif isinstance(it, go.Figure):
                corpo.append(it.to_html(full_html=False, include_plotlyjs="inline" if primeiro_grafico else False,
                                        config={"displaylogo": False}))
                primeiro_grafico = False
            elif isinstance(it, tuple) and it[0] == "h3":
                corpo.append(f"<h3>{html_mod.escape(it[1])}</h3>")
            else:
                corpo.append(f'<p class="nota">{html_mod.escape(str(it))}</p>')
        partes.append(f"<section><h2>{html_mod.escape(nome)}</h2>{''.join(corpo)}</section>")
    return (f'<!doctype html><html lang="pt-BR"><head><meta charset="utf-8">'
            f'<meta name="viewport" content="width=device-width, initial-scale=1"><title>{html_mod.escape(titulo)}</title>'
            f"<style>{CSS}</style></head><body><main><h1>{html_mod.escape(titulo)}</h1>"
            f'<p class="meta">{html_mod.escape(subtitulo)}</p>{"".join(partes)}</main></body></html>')
